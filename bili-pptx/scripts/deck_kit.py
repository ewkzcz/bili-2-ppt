#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""学习笔记 deck 的绘图工具：按调用方给定的样式画形状、写文字、放图、排代码，与任何模版无关。

配色、字体、字号、圆角、投影这些设计取值由调用方读模版原件后给出，本模块不内置任何取值。

文字样式（字典）：
    font        字体名，或 {"latin": 西文字体, "ea": 东亚字体}
    size        字号 pt；bold / italic 布尔；color "#RRGGBB"
    spc         字距，与 pptx 里 a:rPr 的 spc 同单位（1/100 pt）
    highlight   荧光笔底色 "#RRGGBB"
    caps        True 时转大写
    align       left / center / right；valign top / middle / bottom
    spacing     行距倍数；space_after 段后 pt
    emph        **短语** 的覆盖样式（默认只加粗），如 {"bold": True, "color": "#1558E8"}

形状样式（字典）：
    fill "#RRGGBB"、alpha 不透明度%、gradient {"type": "linear"/"radial", "angle", "center",
    "stops": [[位置%, 颜色, 不透明度%], ...]}、line / line_pt / line_alpha / dash、
    radius 圆角英寸（"full" 为胶囊或圆）、shadow {"blur", "dist", "dir", "color", "alpha"}、rotation

所有坐标、尺寸单位为英寸，画布 13.333 × 7.5。anim=N 把动画组号写进形状名，
建完后由 inject_animations.py 注入入场动画并洗掉标记。
"""

from __future__ import annotations

import io
import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = 13.333
SLIDE_H = 7.5
# 模版原件目录：每套模版是同名的一对 <模版名>.pptx 与 <模版名>.html
TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "references"

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# ── 文档与页面 ────────────────────────────────────────────────────────

def new_deck(*, title: str, template: str) -> Presentation:
    """建一份 16:9 的空 deck；template 记进文档属性，校验时按它核对配色。"""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = title
    prs.core_properties.category = template
    return prs


def blank(prs):
    """加一张空白页。"""
    return prs.slides.add_slide(prs.slide_layouts[6])


def background(slide, *, color: str | None = None, image=None) -> None:
    """页面背景：纯色，或一张铺满的背景图（路径或字节；放进页面背景，不占形状）。"""
    if image is None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(_hex(color))
        return
    source = io.BytesIO(image) if isinstance(image, (bytes, bytearray)) else str(image)
    _, rel_id = slide.part.get_or_add_image_part(source)
    c_sld = slide._element.find(qn("p:cSld"))
    existing = c_sld.find(qn("p:bg"))
    if existing is not None:
        c_sld.remove(existing)
    c_sld.insert(0, parse_xml(
        f'<p:bg {nsdecls("p", "a", "r")}><p:bgPr><a:blipFill dpi="0" rotWithShape="1">'
        f'<a:blip r:embed="{rel_id}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch>'
        f'</a:blipFill><a:effectLst/></p:bgPr></p:bg>'))


def template_file(name: str, suffix: str = ".pptx") -> Path:
    """references/ 下某套模版的原件路径（<模版名>.pptx 或 <模版名>.html）。"""
    path = TEMPLATES_DIR / f"{name}{suffix}"
    if not path.is_file():
        raise FileNotFoundError(f"references/ 下没有模版原件 {path.name}")
    return path


# ── 文字度量 ──────────────────────────────────────────────────────────

def _is_mono(style: dict) -> bool:
    font = style.get("font", "")
    name = font.get("latin", "") if isinstance(font, dict) else font
    return any(key in name.lower() for key in ("mono", "consolas", "menlo", "courier", "code"))


def em_width(text: str, style: dict | None = None) -> float:
    """一行文字的排版宽度（em）：中文 1；等宽字体拉丁 0.6；其余大写 0.68、小写 0.55、空格 0.3。"""
    mono = _is_mono(style or {})
    width = 0.0
    for ch in str(text).replace("**", ""):
        if ord(ch) > 0x2E80:
            width += 1.0
        elif mono:
            width += 0.6
        elif ch == " ":
            width += 0.3
        else:
            width += 0.68 if ch.isupper() else 0.55
    spc = (style or {}).get("spc", 0) / 100
    return width + len(str(text)) * spc / max((style or {}).get("size", 14), 1)


def text_width(text: str, style: dict) -> float:
    """一行文字按样式排出来的宽度（英寸）。"""
    return em_width(text, style) * style.get("size", 14) / 72


def text_lines(text, style: dict, width: float) -> int:
    """按宽度估算文字会排成几行（换行符分段，每段至少一行）。"""
    per_line = max(width * 72 / style.get("size", 14), 1.0)
    paragraphs = text if isinstance(text, (list, tuple)) else str(text).split("\n")
    return sum(max(1, math.ceil(em_width(p, style) / per_line)) for p in paragraphs)


def text_height(text, style: dict, width: float) -> float:
    """按宽度估算文字排出来的高度（英寸）。"""
    line_h = style.get("size", 14) * 1.2 * style.get("spacing", 1.0) / 72
    return text_lines(text, style, width) * line_h


def fit_size(text: str, style: dict, width: float) -> float:
    """一行放不下时，把字号收到刚好放下的大小。"""
    need = em_width(text, style)
    return min(style["size"], int(width * 72 / max(need, 0.01)))


# ── 文字 ──────────────────────────────────────────────────────────────

def _hex(value) -> str:
    return str(value).lstrip("#").upper()


def _name_shape(shape, anim) -> None:
    """把动画组号写进形状名（anim-1、anim-2:wipe），供 inject_animations.py 读取。"""
    if anim is not None:
        shape.name = f"anim-{anim}"


def _set_run(run, style: dict) -> None:
    """按文字样式设置一个文字片段：字号、粗斜体、颜色、字距、高亮与西文 / 东亚字体槽。"""
    font = style.get("font", "Arial")
    latin, ea = (font["latin"], font.get("ea", font["latin"])) if isinstance(font, dict) \
        else (font, font)
    run.font.name = latin
    if style.get("size"):
        run.font.size = Pt(style["size"])
    run.font.bold = bool(style.get("bold"))
    run.font.italic = bool(style.get("italic"))
    if style.get("color"):
        run.font.color.rgb = RGBColor.from_string(_hex(style["color"]))
    r_pr = run._r.get_or_add_rPr()
    if style.get("spc"):
        r_pr.set("spc", str(int(style["spc"])))
    latin_el = r_pr.find(qn("a:latin"))
    # rPr 子元素顺序由 schema 规定：latin → ea → cs，从后往前插正好排好
    for tag, face in (("a:cs", latin), ("a:ea", ea)):
        existing = r_pr.find(qn(tag))
        if existing is not None:
            r_pr.remove(existing)
        latin_el.addnext(r_pr.makeelement(qn(tag), {"typeface": face}))
    if style.get("highlight"):
        # highlight 在 schema 里排在填充之后、字体之前
        latin_el.addprevious(parse_xml(
            f'<a:highlight {nsdecls("a")}><a:srgbClr val="{_hex(style["highlight"])}"/>'
            f'</a:highlight>'))


def write_runs(paragraph, text: str, style: dict) -> None:
    """把一段文字写进段落；**短语** 按样式里的 emph 强调（默认只加粗）。"""
    emph = {"bold": True, **style.get("emph", {})}
    content = str(text).upper() if style.get("caps") else str(text)
    for index, part in enumerate(content.split("**")):
        if part:
            run = paragraph.add_run()
            run.text = part
            _set_run(run, {**style, **emph} if index % 2 else style)


def _format_paragraph(paragraph, style: dict) -> None:
    paragraph.alignment = _ALIGN[style.get("align", "left")]
    if style.get("spacing"):
        paragraph.line_spacing = style["spacing"]
    if style.get("space_after"):
        paragraph.space_after = Pt(style["space_after"])


def textbox(slide, text, x, y, w, h, style: dict, *, wrap=True, anim=None, margin=0.0):
    """放一个文本框并返回它。

    text 可以是字符串（换行符分段），也可以是段落列表；段落是字符串，或 [(文字, 样式), ...]
    表示同一段里混排几种样式（样式与 style 合并）。
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.05)),
                                   Inches(max(h, 0.05)))
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = \
        Inches(margin)
    frame.vertical_anchor = _ANCHOR[style.get("valign", "top")]
    paragraphs = text if isinstance(text, (list, tuple)) else str(text).split("\n")
    for index, content in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _format_paragraph(paragraph, style)
        if isinstance(content, (list, tuple)):
            for part, part_style in content:
                write_runs(paragraph, part, {**style, **part_style})
        else:
            write_runs(paragraph, content, style)
    _name_shape(box, anim)
    return box


def bullets(slide, items, x, y, w, h, style: dict, *, bullet: dict, term: dict | None = None,
            anim=None):
    """一列要点。bullet 定项目符号 {"char", "color", "scale", "font", "indent"(em)}；

    term 给出时，「术语：解释」的术语部分按 term 覆盖样式写。
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    size = style.get("size", 14)
    indent = int(Pt(size * bullet.get("indent", 1.2)))
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _format_paragraph(paragraph, style)
        paragraph.space_after = Pt(style.get("space_after", size * 0.6))
        p_pr = paragraph._p.get_or_add_pPr()
        p_pr.set("marL", str(indent))
        p_pr.set("indent", str(-indent))
        clr = p_pr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"),
                                   {"val": _hex(bullet.get("color", style.get("color")))}))
        p_pr.append(clr)
        p_pr.append(p_pr.makeelement(qn("a:buSzPct"),
                                     {"val": str(int(bullet.get("scale", 1.0) * 100000))}))
        if bullet.get("font"):
            p_pr.append(p_pr.makeelement(qn("a:buFont"), {"typeface": bullet["font"]}))
        p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": bullet.get("char", "•")}))
        head, colon, tail = item.partition("：")
        if term and colon and 0 < len(head) <= 16 and "**" not in head:
            write_runs(paragraph, head + colon, {**style, "bold": True, **term})
            item = tail
        write_runs(paragraph, item, style)
    _name_shape(box, anim)
    return box


# ── 形状 ──────────────────────────────────────────────────────────────

def _alpha(parent, alpha) -> None:
    srgb = parent.find(".//" + qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))}))


def _gradient(target, gradient: dict) -> None:
    """线性或径向渐变填充；径向可用 center [x%, y%] 偏移高光位置。"""
    stops = "".join(
        f'<a:gs pos="{int(pos * 1000)}"><a:srgbClr val="{_hex(value)}">'
        f'<a:alpha val="{int((rest[0] if rest else 100) * 1000)}"/></a:srgbClr></a:gs>'
        for pos, value, *rest in gradient["stops"])
    if gradient.get("type") == "radial":
        cx, cy = gradient.get("center", [50, 50])
        shade = (f'<a:path path="circle"><a:fillToRect l="{int(cx * 1000)}" t="{int(cy * 1000)}" '
                 f'r="{int((100 - cx) * 1000)}" b="{int((100 - cy) * 1000)}"/></a:path>')
    else:
        shade = f'<a:lin ang="{int(gradient.get("angle", 0) * 60000)}" scaled="0"/>'
    grad = parse_xml(f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>{stops}</a:gsLst>'
                     f'{shade}</a:gradFill>')
    sp_pr = target.fill._xPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill"):
        for element in sp_pr.findall(qn(tag)):
            sp_pr.remove(element)
    geometry = sp_pr.find(qn("a:prstGeom"))
    (geometry if geometry is not None else sp_pr).addnext(grad)


def _shadow(target, spec: dict) -> None:
    """外投影：blur / dist 英寸，dir 角度，alpha 不透明度%。"""
    sp_pr = target._element.spPr
    effect = sp_pr.makeelement(qn("a:effectLst"), {})
    outer = effect.makeelement(qn("a:outerShdw"), {
        "blurRad": str(Inches(spec.get("blur", 0))), "dist": str(Inches(spec.get("dist", 0))),
        "dir": str(int(spec.get("dir", 90) * 60000)), "rotWithShape": "0"})
    clr = outer.makeelement(qn("a:srgbClr"), {"val": _hex(spec.get("color", "#000000"))})
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(spec.get("alpha", 30) * 1000))}))
    outer.append(clr)
    effect.append(outer)
    sp_pr.append(effect)


def _line_style(target, style: dict) -> None:
    if not style.get("line"):
        target.line.fill.background()
        return
    target.line.color.rgb = RGBColor.from_string(_hex(style["line"]))
    target.line.width = Pt(style.get("line_pt", 1.0))
    if style.get("line_alpha") is not None:
        _alpha(target.line._get_or_add_ln().find(qn("a:solidFill")), style["line_alpha"])
    if style.get("dash"):
        target.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def shape(slide, x, y, w, h, style: dict, *, kind="rect", anim=None):
    """按形状样式画矩形 / 圆角矩形 / 椭圆并返回它。"""
    if kind == "ellipse":
        prst = MSO_SHAPE.OVAL
    else:
        prst = MSO_SHAPE.ROUNDED_RECTANGLE if style.get("radius") else MSO_SHAPE.RECTANGLE
    created = slide.shapes.add_shape(prst, Inches(x), Inches(y), Inches(max(w, 0.001)),
                                     Inches(max(h, 0.001)))
    # 摘掉自动图形自带的主题样式引用，否则主题会把投影、填充带回来
    created.shadow.inherit = False
    theme_style = created._element.find(qn("p:style"))
    if theme_style is not None:
        created._element.remove(theme_style)
    radius = style.get("radius")
    if radius and created.adjustments and len(created.adjustments):
        radius = min(w, h) / 2 if radius == "full" else radius
        created.adjustments[0] = min(radius / max(min(w, h), 0.01), 0.5)
    if style.get("gradient"):
        _gradient(created, style["gradient"])
    elif style.get("fill"):
        created.fill.solid()
        created.fill.fore_color.rgb = RGBColor.from_string(_hex(style["fill"]))
        if style.get("alpha") is not None:
            _alpha(created.fill._xPr.find(qn("a:solidFill")), style["alpha"])
    else:
        created.fill.background()
    _line_style(created, style)
    if style.get("shadow"):
        _shadow(created, style["shadow"])
    if style.get("rotation"):
        created.rotation = style["rotation"]
    _name_shape(created, anim)
    return created


def line(slide, x1, y1, x2, y2, *, color: str, pt: float = 1.0, tail=False, head=False,
         dash=False, anim=None):
    """一根直线；tail / head 在终点 / 起点画实心箭头。"""
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                           Inches(x2), Inches(y2))
    connector.line.color.rgb = RGBColor.from_string(_hex(color))
    connector.line.width = Pt(pt)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = connector.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if tail:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    _name_shape(connector, anim)
    return connector


def picture(slide, image, x, y, w, h, *, frame: dict | None = None, anim=None):
    """把图片等比缩放进给定区域并居中（不拉伸）；frame 给出描边 / 投影。"""
    from PIL import Image

    source = io.BytesIO(image) if isinstance(image, (bytes, bytearray)) else str(image)
    with Image.open(source) as handle:
        image_w, image_h = handle.size
    if hasattr(source, "seek"):
        source.seek(0)
    scale = min(w / image_w, h / image_h)
    draw_w, draw_h = Emu(int(image_w * scale * 914400)), Emu(int(image_h * scale * 914400))
    left = Emu(int(Inches(x)) + (Inches(w) - draw_w) // 2)
    top = Emu(int(Inches(y)) + (Inches(h) - draw_h) // 2)
    pic = slide.shapes.add_picture(source, left, top, draw_w, draw_h)
    if frame:
        _line_style(pic, frame)
        if frame.get("shadow"):
            _shadow(pic, frame["shadow"])
    _name_shape(pic, anim)
    return pic


# ── 代码 ──────────────────────────────────────────────────────────────

_CODE_KEYWORDS = {
    "const", "let", "var", "function", "async", "await", "return", "if", "else", "for",
    "while", "of", "in", "new", "class", "import", "from", "export", "try", "catch",
    "throw", "true", "false", "null", "undefined", "def", "print", "pip", "curl", "npm",
    "node", "break", "continue", "typeof", "this", "with", "as", "None", "True", "False",
}
_CODE_TOKEN = re.compile(r"(\"[^\"]*\"|'[^']*'|`[^`]*`)|(\b\d+(?:\.\d+)?\b)|([A-Za-z_]\w*)")


def _code_runs(paragraph, text: str, style: dict, syntax: dict) -> None:
    """写一行代码：注释、关键字、字符串、数值按 syntax 着色，其余用 style 的颜色。"""
    def put(part, role=None):
        if part or not paragraph.runs:
            run = paragraph.add_run()
            run.text = part
            _set_run(run, {**style, "color": syntax.get(role, style["color"]) if role
                           else style["color"]})

    stripped = text.lstrip()
    if stripped.startswith(("//", "# ", "--")) or stripped == "#":
        put(text, "comment")
        return
    comment = ""
    match = re.search(r"(?<![:\w\"'])//", text)
    if match:
        text, comment = text[:match.start()], text[match.start():]
    pos, plain = 0, ""
    for token in _CODE_TOKEN.finditer(text):
        plain += text[pos:token.start()]
        pos = token.end()
        if token.group(1):
            role = "string"
        elif token.group(2):
            role = "number" if "number" in syntax else "string"
        else:
            role = "keyword" if token.group(3) in _CODE_KEYWORDS else None
        if role and syntax.get(role):
            put(plain)
            plain = ""
            put(token.group(0), role)
        else:
            plain += token.group(0)
    put(plain + text[pos:])
    if comment:
        put(comment, "comment")


def code_block(slide, code: str, x, y, w, h, *, box: dict, style: dict, syntax: dict,
               pad=(0.2, 0.17), anim=None):
    """等宽代码块：底框按 box 样式，字按 style，关键字 / 字符串 / 数值 / 注释按 syntax 着色。"""
    frame_shape = shape(slide, x, y, w, h, box, anim=anim)
    frame = frame_shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(pad[0])
    frame.margin_top = frame.margin_bottom = Inches(pad[1])
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, text in enumerate(code.splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = style.get("spacing", 1.18)
        paragraph.alignment = PP_ALIGN.LEFT
        _code_runs(paragraph, text, style, syntax)
    return frame_shape


# ── 存盘 ──────────────────────────────────────────────────────────────

def save(prs: Presentation, path, *, heading_font: dict, body_font: dict) -> Path:
    """写盘，并把 pptx 主题的标题 / 正文字体设成给定字体（续编时新文本框也是这套字体）。

    字体取 {"latin": ..., "ea": ...}。
    """
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for tag, faces in (("majorFont", heading_font), ("minorFont", body_font)):
        font = root.find(f".//a:{tag}", ns)
        font.find("a:latin", ns).set("typeface", faces["latin"])
        font.find("a:ea", ns).set("typeface", faces.get("ea", faces["latin"]))
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
