#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 pptx 渲染成逐页图片，用来做视觉自检。

只看 XML 校验通过是不够的：文字压线、图片出框、两栏对不齐这些只有看图才发现。
流程是 pptx → PDF（LibreOffice）→ PNG（poppler）。

**动画不会出现在渲染结果里。** 渲染看到的是所有元素都可见的完整状态——
PowerPoint 的入场动画只控制「什么时候显示」，元素本身始终在页面上。
这正好是我们要的：静态看是一页完整的图，放映时点一下多一块。
所以如果渲染出来有重叠，放映时同样会重叠，必须在这一步修掉。

用法：
    python3 render_preview.py deck.pptx --out-dir preview/
    python3 render_preview.py deck.pptx -o preview/ --slides 3,7,12

自检只渲染需要复查的页（建页前评估的高风险页 + validate_deck.py --suspects 报出的页），
不做全量逐页扫描；全量渲染只在看模版原件时用。
带 --slides 时先复制一份只含这几页的 pptx 再交给 LibreOffice，几百页的 deck 也只转要看的页；
输出图片仍按原页号命名（slide-07.jpg 就是原 deck 第 7 页）。页码域会显示成裁剪后的序号，看版面时忽略即可。
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


# LibreOffice 在不同装法下的位置。macOS 用 brew 装的 cask 在 /Applications，
# Linux 和 CI 镜像里通常直接在 PATH 上。
SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
    "/opt/homebrew/bin/soffice",
    "/snap/bin/libreoffice",
)

# fontconfig 配置目录候选：LibreOffice 自带的配置不扫描 macOS 系统字体目录，
# 中文会渲染成空白，改用能读到系统字体的配置目录。
FONTCONFIG_CANDIDATES = (
    "/opt/homebrew/etc/fonts",
    "/usr/local/etc/fonts",
    "/etc/fonts",
)


def build_env() -> dict:
    """把 fontconfig 指向一份能看见系统中文字体的配置。"""
    env = dict(os.environ)
    if env.get("FONTCONFIG_PATH"):
        return env
    for candidate in FONTCONFIG_CANDIDATES:
        if (Path(candidate) / "fonts.conf").is_file():
            env["FONTCONFIG_PATH"] = candidate
            break
    return env


def find_soffice() -> str:
    """找一个能用的 LibreOffice 可执行文件。"""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in SOFFICE_CANDIDATES:
        if Path(candidate).is_file():
            return candidate
    raise SystemExit(
        "找不到 LibreOffice。装一个再跑：brew install --cask libreoffice\n"
        "（没有它就没法把 pptx 渲染成图，视觉自检这一步做不了）"
    )


def find_pdftoppm() -> str:
    """找 poppler 的 pdftoppm。"""
    found = shutil.which("pdftoppm")
    if found:
        return found
    for candidate in ("/opt/homebrew/bin/pdftoppm", "/usr/local/bin/pdftoppm"):
        if Path(candidate).is_file():
            return candidate
    raise SystemExit("找不到 pdftoppm。装一个再跑：brew install poppler")


def to_pdf(pptx: Path, out_dir: Path) -> Path:
    """用 LibreOffice 把 pptx 转成 PDF。"""
    soffice = find_soffice()
    out_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(pptx.resolve()),
        ],
        capture_output=True,
        text=True,
        timeout=300,
        env=build_env(),
    )
    pdf = out_dir / f"{pptx.stem}.pdf"
    if result.returncode != 0 or not pdf.is_file():
        raise SystemExit(
            f"LibreOffice 转换失败（退出码 {result.returncode}）\n"
            f"stdout: {result.stdout.strip()}\nstderr: {result.stderr.strip()}"
        )
    return pdf


def trim_deck(pptx: Path, slides: list[int], work_dir: Path) -> tuple[Path, int] | None:
    """复制一份只保留指定页的 pptx，返回（副本路径，原 deck 总页数）；页号全覆盖或读不了时返回 None。"""
    try:
        from pptx import Presentation
    except ImportError:
        return None
    prs = Presentation(str(pptx))
    id_list = prs.slides._sldIdLst
    total = len(id_list)
    keep = {no for no in slides if 1 <= no <= total}
    if not keep or len(keep) == total:
        return None
    for index, sld_id in enumerate(list(id_list), start=1):
        if index not in keep:
            prs.part.drop_rel(sld_id.rId)
            id_list.remove(sld_id)
    trimmed = work_dir / pptx.name
    prs.save(str(trimmed))
    return trimmed, total


def to_images(pdf: Path, out_dir: Path, *, dpi: int, slides: list[int] | None,
              names: list[int] | None = None, width: int = 0) -> list[Path]:
    """把 PDF 转成 JPEG：slides 为空时全部页，否则只转指定的几页。

    names 给出时（裁剪副本），第 i 个 PDF 页存成 slide-<names[i]>.jpg，按原页号命名。
    """
    pdftoppm = find_pdftoppm()
    out_dir.mkdir(parents=True, exist_ok=True)
    if names:
        jobs = [["-f", str(page), "-l", str(page), "-singlefile", str(pdf), str(out_dir / f"slide-{no:0{width}d}")]
                for page, no in enumerate(names, start=1)]
    else:
        ranges = [["-f", str(no), "-l", str(no)] for no in slides] if slides else [[]]
        jobs = [[*r, str(pdf), str(out_dir / "slide")] for r in ranges]
    for job in jobs:
        result = subprocess.run([pdftoppm, "-jpeg", "-r", str(dpi), *job], capture_output=True, text=True, timeout=300)
        if result.returncode != 0:
            raise SystemExit(
                f"pdftoppm 失败（退出码 {result.returncode}）\n{result.stderr.strip()}"
            )
    if names:
        return [out_dir / f"slide-{no:0{width}d}.jpg" for no in names]
    return sorted(out_dir.glob("slide-*.jpg"))


def main() -> int:
    parser = argparse.ArgumentParser(description="把 pptx 渲染成逐页图片供视觉自检")
    parser.add_argument("pptx", type=Path, help="输入 pptx")
    parser.add_argument("--out-dir", "-o", type=Path, default=None, help="输出目录")
    parser.add_argument("--dpi", type=int, default=110, help="渲染分辨率（默认 110）")
    parser.add_argument("--slide", type=int, default=None, help="只渲染某一页")
    parser.add_argument("--slides", default=None, help="只渲染这几页，逗号分隔，如 3,7,12")
    parser.add_argument(
        "--keep-pdf", action="store_true", help="保留中间 PDF，方便直接翻看"
    )
    args = parser.parse_args()

    if not args.pptx.is_file():
        raise SystemExit(f"找不到文件：{args.pptx}")

    out_dir = args.out_dir or (args.pptx.parent / f"{args.pptx.stem}-preview")
    out_dir.mkdir(parents=True, exist_ok=True)

    slides = [int(no) for no in args.slides.split(",") if no.strip()] if args.slides else []
    if args.slide is not None:
        slides.append(args.slide)
    slides = sorted(set(slides))

    with tempfile.TemporaryDirectory(prefix="render_preview_") as tmp:
        trimmed = trim_deck(args.pptx, slides, Path(tmp)) if slides else None
        if trimmed:
            deck, total = trimmed
            pdf = to_pdf(deck, Path(tmp))
            kept = [no for no in slides if 1 <= no <= total]
            images = to_images(pdf, out_dir, dpi=args.dpi, slides=None, names=kept, width=len(str(total)))
        else:
            pdf = to_pdf(args.pptx, out_dir)
            images = to_images(pdf, out_dir, dpi=args.dpi, slides=slides or None)
        if args.keep_pdf and pdf.parent != out_dir:
            shutil.copy2(pdf, out_dir / pdf.name)
        elif not args.keep_pdf:
            pdf.unlink(missing_ok=True)

    if not images:
        print("没有生成任何图片，检查一下这份 pptx 是不是空的", file=sys.stderr)
        return 1

    print(f"渲染出 {len(images)} 页，在 {out_dir}")
    for image in images:
        print(f"  {image}")
    print()
    print("只看重叠、溢出、截断、出框这类主要问题，修完重渲染改过的页即可。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
