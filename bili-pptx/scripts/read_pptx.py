#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把一份 pptx 的设计元数据整理成可读清单：主题配色与字体、每页背景、每个形状的位置与样式、每段文字的字样。

读模版原件时用它代替逐个翻 XML：数值与 XML 一致（坐标英寸、字号 pt、字距 spc 为 1/100 pt、
圆角换算成英寸），看完再按内容自己构图。最后列出「特殊样式清单」（文字描边、渐变字、固定行距、
箭头、表格……）和出现的页，样式表要把清单里的每一项都写进去；同名 html 里的 CSS 特效也一并列出。

用法：
    read_pptx.py <原件.pptx>                 # 全部页
    read_pptx.py <原件.pptx> --slides 1,3,4  # 只列这些页（主题、用量与特殊样式清单照常汇总全篇）
"""

from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from pptx_effects import EFFECTS, effects_of, how_to

EMU = 914400

# 同名 html 里值得留意的 CSS 特效：pptx 里要找到对应写法，找不到的在样式表里说明取舍
CSS_EFFECTS = (
    ("文字描边", r"text-stroke"),
    ("渐变文字", r"background-clip:\s*text"),
    ("文字投影", r"text-shadow"),
    ("毛玻璃", r"backdrop-filter"),
    ("投影", r"box-shadow"),
    ("线性渐变", r"linear-gradient"),
    ("径向渐变", r"radial-gradient"),
    ("锥形渐变", r"conic-gradient"),
    ("滤镜", r"(?<![-\w])filter:"),
    ("混合模式", r"mix-blend-mode"),
    ("裁切形状", r"clip-path"),
    ("虚线边框", r"dashed"),
    ("斜体", r"font-style:\s*italic"),
    ("竖排", r"writing-mode"),
)


def _color(parent) -> str:
    """一个填充 / 描边节点下的颜色与不透明度，如 1558E8 或 FFFFFF@62%。"""
    if parent is None:
        return ""
    clr = parent.find(".//" + qn("a:srgbClr"))
    if clr is None:
        scheme = parent.find(".//" + qn("a:schemeClr"))
        return f"scheme:{scheme.get('val')}" if scheme is not None else ""
    alpha = clr.find(qn("a:alpha"))
    return clr.get("val") + (f"@{int(alpha.get('val')) // 1000}%" if alpha is not None else "")


def _gradient(grad) -> str:
    stops = " ".join(f"{int(gs.get('pos')) // 1000}%:{_color(gs)}" for gs in grad.iter(qn("a:gs")))
    lin = grad.find(qn("a:lin"))
    kind = f"lin{int(lin.get('ang', 0)) // 60000}°" if lin is not None else "radial"
    return f"gradient({kind} {stops})"


def _line(ln, prefix="line") -> str:
    """描边：颜色 / 粗细、虚线、箭头。"""
    if ln is None or ln.find(qn("a:noFill")) is not None:
        return ""
    color = _color(ln.find(qn("a:solidFill")))
    if not color:
        return ""
    text = f"{prefix}={color}/{int(ln.get('w', '12700')) / 12700:g}pt"
    dash = ln.find(qn("a:prstDash"))
    if dash is not None and dash.get("val") != "solid":
        text += f" dash={dash.get('val')}"
    for tag, label in (("a:headEnd", "head"), ("a:tailEnd", "tail")):
        end = ln.find(qn(tag))
        if end is not None and end.get("type", "none") != "none":
            text += f" {label}={end.get('type')}"
    return text


def _shape_style(shape) -> str:
    """形状的填充、渐变、描边、圆角、投影、旋转翻转摘要。"""
    sp_pr = shape._element.find(qn("p:spPr"))
    if sp_pr is None:
        return ""
    parts = []
    fill = _color(sp_pr.find(qn("a:solidFill")))
    if fill:
        parts.append(f"fill={fill}")
    grad = sp_pr.find(qn("a:gradFill"))
    if grad is not None:
        parts.append(_gradient(grad))
    line = _line(sp_pr.find(qn("a:ln")))
    if line:
        parts.append(line)
    shadow = sp_pr.find(".//" + qn("a:outerShdw"))
    if shadow is not None:
        parts.append(f"shadow(blur={int(shadow.get('blurRad', 0)) / EMU:.3f} "
                     f"dist={int(shadow.get('dist', 0)) / EMU:.3f} "
                     f"dir={int(shadow.get('dir', 0)) // 60000} {_color(shadow)})")
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is not None and geom.get("prst") not in (None, "rect"):
        adj = geom.find(".//" + qn("a:gd"))
        if geom.get("prst") == "roundRect":
            ratio = int(adj.get("fmla").split()[-1]) / 100000 if adj is not None else 0.16667
            parts.append(f"roundRect(r={min(shape.width, shape.height) / EMU * ratio:.3f})")
        else:
            parts.append(geom.get("prst"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is not None:
        if xfrm.get("rot") not in (None, "0"):
            parts.append(f"rot={int(xfrm.get('rot')) / 60000:g}°")
        parts += [flag for flag in ("flipH", "flipV") if xfrm.get(flag) == "1"]
    return " ".join(parts)


def _run_style(r_pr) -> str:
    """一个文字片段的字样：字号、粗斜体、颜色 / 渐变、西文 / 东亚字体、字距、描边、高亮。"""
    if r_pr is None:
        return "(继承)"
    latin, ea = r_pr.find(qn("a:latin")), r_pr.find(qn("a:ea"))
    highlight = r_pr.find(qn("a:highlight"))
    grad = r_pr.find(qn("a:gradFill"))
    fill = _gradient(grad) if grad is not None else _color(r_pr.find(qn("a:solidFill"))) or "-"
    outline = _line(r_pr.find(qn("a:ln")), "outline")
    flags = ("b" if r_pr.get("b") == "1" else "") + ("i" if r_pr.get("i") == "1" else "")
    size = f"{int(r_pr.get('sz')) / 100:g}pt" if r_pr.get("sz") else "继承字号"
    return (f"{size}{' ' + flags if flags else ''} {fill} "
            f"{latin.get('typeface') if latin is not None else '-'}/"
            f"{ea.get('typeface') if ea is not None else '-'}"
            f"{' spc=' + r_pr.get('spc') if r_pr.get('spc') else ''}"
            f"{' ' + outline if outline else ''}"
            f"{' highlight=' + _color(highlight) if highlight is not None else ''}")


def _paragraph_marks(paragraph) -> str:
    """段落的对齐、缩进、行距（倍数或固定 pt）、段前段后距与项目符号 / 编号。"""
    p_pr = paragraph.find(qn("a:pPr"))
    if p_pr is None:
        return ""
    parts = []
    if p_pr.get("algn"):
        parts.append(f"algn={p_pr.get('algn')}")
    if p_pr.get("marL") not in (None, "0") or p_pr.get("indent") not in (None, "0"):
        parts.append(f"marL={int(p_pr.get('marL', 0)) / EMU:.3f} "
                     f"indent={int(p_pr.get('indent', 0)) / EMU:.3f}")
    for tag, label in (("a:lnSpc", "lnSpc"), ("a:spcBef", "spcBef"), ("a:spcAft", "spcAft")):
        spacing = p_pr.find(qn(tag))
        if spacing is None:
            continue
        pct, pts = spacing.find(qn("a:spcPct")), spacing.find(qn("a:spcPts"))
        if pct is not None:
            parts.append(f"{label}={int(pct.get('val')) / 100000:g}倍")
        elif pts is not None:
            parts.append(f"{label}={int(pts.get('val')) / 100:g}pt")
    bullet, auto = p_pr.find(qn("a:buChar")), p_pr.find(qn("a:buAutoNum"))
    if bullet is not None or auto is not None:
        mark = f"bullet={bullet.get('char')}" if bullet is not None \
            else f"autonum={auto.get('type')}" + (f"@{auto.get('startAt')}" if auto.get("startAt") else "")
        size = p_pr.find(qn("a:buSzPct"))
        font = p_pr.find(qn("a:buFont"))
        parts.append(" ".join(filter(None, (
            mark, _color(p_pr.find(qn("a:buClr"))),
            f"size={int(size.get('val')) / 1000:g}%" if size is not None else "",
            f"font={font.get('typeface')}" if font is not None else ""))))
    return " ".join(parts)


def _print_text(tx_body, indent: str) -> None:
    """一个文本体的换行 / 内边距 / 锚点，以及每段每个片段的字样。"""
    body = tx_body.find(qn("a:bodyPr"))
    if body is not None:
        insets = [body.get(k) for k in ("lIns", "tIns", "rIns", "bIns")]
        marks = " ".join(filter(None, (
            f"anchor={body.get('anchor')}" if body.get("anchor") else "",
            "wrap=none" if body.get("wrap") == "none" else "",
            "inset=" + "/".join(f"{int(v) / EMU:.3f}" for v in insets) if any(insets) else "",
            "autofit" if body.find(qn("a:spAutoFit")) is not None else "")))
        if marks:
            print(f"{indent}{marks}")
    for paragraph in tx_body.iter(qn("a:p")):
        marks = _paragraph_marks(paragraph)
        for run in paragraph.findall(qn("a:r")):
            text = run.findtext(qn("a:t")) or ""
            if text.strip():
                print(f"{indent}「{text[:40]}」 {_run_style(run.find(qn('a:rPr')))} {marks}".rstrip())


def _print_table(shape, indent: str) -> None:
    """表格：列宽、行高，每个单元格的底色、四边框线、内边距、锚点与文字。"""
    table = shape.table
    print(f"{indent}列宽 {' '.join(f'{c.width / EMU:.3f}' for c in table.columns)}  "
          f"行高 {' '.join(f'{r.height / EMU:.3f}' for r in table.rows)}")
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            tc_pr = cell._tc.find(qn("a:tcPr"))
            marks = []
            if tc_pr is not None:
                fill = _color(tc_pr.find(qn("a:solidFill")))
                if fill:
                    marks.append(f"fill={fill}")
                for side in ("L", "R", "T", "B"):
                    line = _line(tc_pr.find(qn(f"a:ln{side}")), f"border{side}")
                    if line:
                        marks.append(line)
                pads = [tc_pr.get(k) for k in ("marL", "marT", "marR", "marB")]
                if any(pads):
                    marks.append("pad=" + "/".join(f"{int(v or 0) / EMU:.3f}" for v in pads))
                if tc_pr.get("anchor"):
                    marks.append(f"anchor={tc_pr.get('anchor')}")
            print(f"{indent}[{r},{c}] {' '.join(marks)}")
            _print_text(cell._tc.find(qn("a:txBody")), indent + "  ")


def _print_shapes(shapes, indent: str) -> None:
    for shape in shapes:
        box = (f"({shape.left / EMU:.3f}, {shape.top / EMU:.3f}) "
               f"{shape.width / EMU:.3f}×{shape.height / EMU:.3f}")
        print(f"{indent}#{shape.shape_id:<4} {str(shape.shape_type):<16} {box:<32} {_shape_style(shape)}")
        if shape.shape_type is not None and shape.shape_type == 6:  # GROUP
            _print_shapes(shape.shapes, indent + "  ")
        elif getattr(shape, "has_table", False) and shape.has_table:
            _print_table(shape, indent + "      ")
        elif shape.has_text_frame and shape.text_frame.text.strip():
            _print_text(shape.text_frame._txBody, indent + "      ")


def _theme(prs) -> None:
    """主题里的配色方案与主 / 次字体。"""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = part._element if hasattr(part, "_element") else None
    if root is None:
        from lxml import etree
        root = etree.fromstring(part.blob)
    scheme = root.find(".//" + qn("a:clrScheme"))
    if scheme is not None:
        colors = []
        for child in scheme:
            value = child.find(qn("a:srgbClr"))
            system = child.find(qn("a:sysClr"))
            colors.append(f"{child.tag.split('}')[1]}="
                          f"{value.get('val') if value is not None else system.get('lastClr')}")
        print("主题配色：", " ".join(colors))
    for tag in ("majorFont", "minorFont"):
        font = root.find(".//" + qn(f"a:{tag}"))
        if font is not None:
            print(f"主题字体 {tag}：latin={font.find(qn('a:latin')).get('typeface')} "
                  f"ea={font.find(qn('a:ea')).get('typeface')}")


def _background(slide) -> str:
    bg = slide._element.find(".//" + qn("p:bgPr"))
    if bg is None:
        return ""
    if bg.find(".//" + qn("a:blip")) is not None:
        return "图片"
    grad = bg.find(qn("a:gradFill"))
    return _gradient(grad) if grad is not None else _color(bg.find(qn("a:solidFill")))


def _css_effects(html: Path) -> None:
    if not html.is_file():
        return
    text = html.read_text(encoding="utf-8", errors="ignore")
    found = [f"{name}×{n}" for name, pattern in CSS_EFFECTS if (n := len(re.findall(pattern, text)))]
    if found:
        print(f"\nhtml 里的 CSS 特效（{html.name}）：", ", ".join(found))


def read(path: Path, slides: set[int] | None) -> None:
    """打印主题、逐页形状清单，并汇总全篇的字体、颜色用量和特殊样式清单。"""
    prs = Presentation(str(path))
    print(f"画布 {prs.slide_width / EMU:.3f} × {prs.slide_height / EMU:.3f} 英寸，共 {len(prs.slides)} 页")
    _theme(prs)
    fonts, colors = Counter(), Counter()
    effects: dict[str, list[int]] = defaultdict(list)
    for number, slide in enumerate(prs.slides, 1):
        for element in slide._element.iter(qn("a:latin"), qn("a:ea")):
            fonts[element.get("typeface")] += 1
        for element in slide._element.iter(qn("a:srgbClr")):
            colors[element.get("val")] += 1
        for name in effects_of(slide._element):
            effects[name].append(number)
        if slides and number not in slides:
            continue
        print(f"\n== 第 {number} 页  背景 {_background(slide) or '继承'}")
        _print_shapes(slide.shapes, "  ")
    print("\n字体用量：", ", ".join(f"{k}×{v}" for k, v in fonts.most_common(12)))
    print("颜色用量：", ", ".join(f"{k}×{v}" for k, v in colors.most_common(30)))
    print("\n特殊样式清单（样式表要把每一项写进去；deck_kit 写法见右）：")
    for name, _, _, _ in EFFECTS:
        if effects.get(name):
            pages = effects[name]
            where = ",".join(map(str, pages[:12])) + ("…" if len(pages) > 12 else "")
            print(f"  {name:<8} 第 {where} 页  → {how_to(name)}")
    _css_effects(path.with_suffix(".html"))


def main() -> int:
    parser = argparse.ArgumentParser(description="列出 pptx 的设计元数据")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--slides", default="", help="只列这些页，逗号分隔")
    args = parser.parse_args()
    wanted = {int(x) for x in args.slides.split(",") if x.strip()} or None
    read(args.pptx, wanted)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
